"""Excel-отчёт за период: итоги по сотрудникам, табель по дням, список по дням."""
from collections import defaultdict
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import WEEKDAYS_SHORT, db, period_title, price

BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def _header(ws, row: int, values: list):
    for col, v in enumerate(values, 1):
        c = ws.cell(row=row, column=col, value=v)
        c.font = BOLD
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_excel(d1: str, d2: str) -> bytes:
    orders = db.period_orders(d1, d2)
    p = price()
    wb = Workbook()

    # --- Итоги ---
    ws = wb.active
    ws.title = "Итоги"
    ws["A1"] = f"Обеды за {period_title(d1, d2)}"
    ws["A1"].font = Font(bold=True, size=14)
    headers = ["№", "Сотрудник", "Дней", "Порций"] + (["Сумма, сум"] if p else [])
    _header(ws, 3, headers)
    stats = db.stats(d1, d2)
    for i, r in enumerate(stats, 1):
        row = [i, r["name"], r["days"], r["portions"]] + ([r["portions"] * p] if p else [])
        ws.append(row)
    total_row = ws.max_row + 1
    ws.cell(row=total_row, column=2, value="Итого").font = BOLD
    ws.cell(row=total_row, column=4, value=sum(r["portions"] for r in stats)).font = BOLD
    if p:
        ws.cell(row=total_row, column=5, value=sum(r["portions"] for r in stats) * p).font = BOLD
        for r in range(4, total_row + 1):
            ws.cell(row=r, column=5).number_format = "# ##0"
    for col, w in zip("ABCDE", (5, 32, 8, 9, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    # --- Табель: сотрудники × дни ---
    ws = wb.create_sheet("Табель")
    days = sorted({o["day"] for o in orders})
    grid: dict[str, dict[str, int]] = defaultdict(dict)
    for o in orders:
        grid[o["name"]][o["day"]] = o["portions"]
    _header(ws, 1, ["Сотрудник"] + [f"{date.fromisoformat(d):%d.%m}\n{WEEKDAYS_SHORT[date.fromisoformat(d).weekday()]}"
                                    for d in days] + ["Итого"])
    for name in sorted(grid, key=str.lower):
        ws.append([name] + [grid[name].get(d) for d in days] + [sum(grid[name].values())])
    footer = ["Порций за день"] + [sum(grid[n].get(d, 0) for n in grid) for d in days] + [sum(o["portions"] for o in orders)]
    ws.append(footer)
    for c in ws[ws.max_row]:
        c.font = BOLD
    ws.column_dimensions["A"].width = 32
    for col in range(2, len(days) + 3):
        ws.column_dimensions[get_column_letter(col)].width = 7
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"

    # --- По дням ---
    ws = wb.create_sheet("По дням")
    _header(ws, 1, ["Дата", "День", "Порций", "Кто обедал"])
    by_day: dict[str, list] = defaultdict(list)
    for o in orders:
        by_day[o["day"]].append(o)
    for d in days:
        rows = by_day[d]
        names = ", ".join(o["name"] + (f" ×{o['portions']}" if o["portions"] > 1 else "") for o in rows)
        dd = date.fromisoformat(d)
        ws.append([f"{dd:%d.%m.%Y}", WEEKDAYS_SHORT[dd.weekday()], sum(o["portions"] for o in rows), names])
    for col, w in zip("ABCD", (12, 6, 9, 100)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
